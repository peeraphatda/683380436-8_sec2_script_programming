# document-automator/src/word_processor.py
import os

import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .utils import setup_logging

logger = setup_logging(__name__)


class WordProcessor:
    """Handles operations related to Word documents using python-docx."""

    def __init__(self):
        logger.info("WordProcessor initialized.")

    def create_document(self):
        """Creates a new blank Word document."""
        document = docx.Document()
        logger.info("New Word document created.")
        return document

    def add_heading(self, document, text, level=1):
        """Adds a heading to the document."""
        document.add_heading(text, level=level)
        logger.info(f"Added heading (Level {level}): '{text}'")

    def add_paragraph(self, document, text, style=None, bold=False, italic=False,
                      font_size=None, align='left'):
        """Adds a paragraph with optional formatting."""
        paragraph = document.add_paragraph(text, style=style)

        # Apply run-level formatting if specified
        if paragraph.runs:
            run = paragraph.runs[0]
            if bold:
                run.bold = True
            if italic:
                run.italic = True
            if font_size:
                run.font.size = Pt(font_size)

        # Apply paragraph alignment
        alignments = {
            'center': WD_ALIGN_PARAGRAPH.CENTER,
            'right': WD_ALIGN_PARAGRAPH.RIGHT,
            'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
        }
        paragraph.alignment = alignments.get(align, WD_ALIGN_PARAGRAPH.LEFT)

        logger.info(f"Added paragraph: '{text[:50]}...'")

    def add_table(self, document, data, include_header=True, header_bold=True, col_widths=None):
        """
        Adds a table to the document.
        `data` is a list of lists (rows). `col_widths` is a list of widths in inches.
        """
        if not data or not data[0]:
            logger.warning("Table data is empty or malformed. Skipping table creation.")
            return

        rows, cols = len(data), len(data[0])
        table = document.add_table(rows=rows, cols=cols)
        table.style = 'Table Grid'

        # Populate table
        for r_idx, row_data in enumerate(data):
            for c_idx, cell_data in enumerate(row_data):
                cell = table.cell(r_idx, c_idx)
                cell.text = str(cell_data)
                if include_header and header_bold and r_idx == 0 and cell.paragraphs[0].runs:
                    cell.paragraphs[0].runs[0].bold = True

        # Set column widths (must be set per cell for Word to respect them)
        if col_widths and len(col_widths) == cols:
            for row in table.rows:
                for c_idx, width in enumerate(col_widths):
                    row.cells[c_idx].width = Inches(width)

        logger.info(f"Added table with {rows} rows and {cols} columns.")

    def add_picture(self, document, image_path, width_inches=None):
        """Adds an image to the document with optional width scaling."""
        if not os.path.exists(image_path):
            logger.error(f"Image file not found: {image_path}")
            return False

        try:
            if width_inches:
                document.add_picture(image_path, width=Inches(width_inches))
            else:
                document.add_picture(image_path)
            logger.info(f"Added picture: {os.path.basename(image_path)}")
            return True
        except Exception as e:
            logger.error(f"Error adding picture '{image_path}': {e}")
            return False

    def save_document(self, document, output_path):
        """Saves the Word document to a specified path."""
        try:
            document.save(output_path)
            logger.info(f"Word document saved to '{output_path}'.")
            return True
        except Exception as e:
            logger.error(f"Error saving Word document to '{output_path}': {e}")
            return False
