# src/spreadsheet_agent.py
import datetime
import json
from copy import copy

import openpyxl
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, FormulaRule, IconSetRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation

from .utils import setup_logging

logger = setup_logging(__name__)

FONT_NAME = "Arial"


class SpreadsheetAgent:
    """
    Agent ที่ทำงานกับ Excel workbook ตามลำดับ "task" ที่กำหนดใน config
    (perception = input file + tasks, action = task handlers)
    """

    def __init__(self, config):
        self.config = config
        self.workbook = None
        self.audit_log = []
        self.start_time = datetime.datetime.now()

        # Task dispatcher: ชื่อ type -> method ที่จัดการ
        self.handlers = {
            "copy_data": self._process_copy_data_task,
            "calculate_column": self._process_calculate_column_task,
            "conditional_format": self._process_conditional_format_task,
            "summary": self._process_summary_task,
            "chart": self._process_chart_task,
            "data_validation": self._process_data_validation_task,
        }
        logger.info("Spreadsheet Agent initialized.")

    # ------------------------------------------------------------------ #
    # Helpers
    # ------------------------------------------------------------------ #
    def _log_audit(self, status, message, task_type="N/A", details=None):
        entry = {
            "timestamp": datetime.datetime.now().isoformat(timespec="seconds"),
            "status": status,
            "task_type": task_type,
            "message": message,
            "details": details or {},
        }
        self.audit_log.append(entry)
        if status in ("ERROR", "CRITICAL"):
            logger.error(f"AUDIT {status}: {message} | Task: {task_type} | Details: {details}")
        elif status == "WARNING":
            logger.warning(f"AUDIT {status}: {message} | Task: {task_type}")
        else:
            logger.info(f"AUDIT {status}: {message} | Task: {task_type}")

    def _load_workbook(self):
        input_file = self.config["input_file"]
        try:
            self.workbook = openpyxl.load_workbook(input_file)
            self._log_audit("SUCCESS", f"Workbook '{input_file}' loaded.", "load_workbook")
            return True
        except FileNotFoundError:
            self._log_audit("CRITICAL", f"Input file not found: {input_file}", "load_workbook")
            return False
        except Exception as e:
            self._log_audit("CRITICAL", f"Failed to load workbook '{input_file}': {e}", "load_workbook")
            return False

    def _get_sheet(self, sheet_name, create=False):
        """คืน worksheet; ถ้าไม่พบ -> สร้างใหม่เมื่อ create=True, ไม่งั้นคืน None"""
        if sheet_name in self.workbook.sheetnames:
            return self.workbook[sheet_name]
        if create:
            self._log_audit("WARNING", f"Sheet '{sheet_name}' not found. Creating it.", "get_sheet")
            return self.workbook.create_sheet(sheet_name)
        return None

    @staticmethod
    def _quote(sheet_name):
        """ใส่ quote ชื่อ sheet สำหรับใช้ใน formula (ชื่อที่มีช่องว่างต้องมี quote)"""
        return "'" + sheet_name.replace("'", "''") + "'"

    # ------------------------------------------------------------------ #
    # Task handlers
    # ------------------------------------------------------------------ #
    def _process_copy_data_task(self, task):
        src_name, dest_name = task["source_sheet"], task["dest_sheet"]
        src = self._get_sheet(src_name)
        if src is None:
            self._log_audit("ERROR", f"Source sheet '{src_name}' not found.", "copy_data", task)
            return
        dest = self._get_sheet(dest_name, create=True)

        count = 0
        for row in src.iter_rows():
            for cell in row:
                new = dest[cell.coordinate]
                new.value = cell.value
                if cell.has_style:
                    new.font = copy(cell.font)
                    new.number_format = cell.number_format
                    new.alignment = copy(cell.alignment)
                    new.fill = copy(cell.fill)
                    new.border = copy(cell.border)
                count += 1
        self._log_audit("SUCCESS", f"Copied {count} cells from '{src_name}' to '{dest_name}'.", "copy_data")

    def _process_calculate_column_task(self, task):
        sheet_name = task["sheet"]
        col = task["target_column"]
        formula = task["formula"]
        start_row = task.get("start_row", 2)
        sheet = self._get_sheet(sheet_name)
        if sheet is None:
            self._log_audit("ERROR", f"Sheet '{sheet_name}' not found.", "calculate_column", task)
            return

        header = sheet[f"{col}1"]
        header.value = task.get("header", "Calculated Column")
        header.font = Font(name=FONT_NAME, bold=True)

        # นับแถวสุดท้ายจากคอลัมน์ A (กันไม่ให้ใส่สูตรลงแถวว่าง)
        last_row = max(
            (r for r in range(start_row, sheet.max_row + 1) if sheet.cell(r, 1).value not in (None, "")),
            default=start_row - 1,
        )
        written = 0
        for r in range(start_row, last_row + 1):
            try:
                cell = sheet[f"{col}{r}"]
                cell.value = formula.replace("{row}", str(r))
                if task.get("number_format"):
                    cell.number_format = task["number_format"]
                written += 1
            except Exception as e:
                self._log_audit("ERROR", f"Failed to write formula at {col}{r}: {e}", "calculate_column", task)
        self._log_audit(
            "SUCCESS", f"Column {col} ('{header.value}') filled in '{sheet_name}' ({written} rows).", "calculate_column"
        )

    def _process_conditional_format_task(self, task):
        sheet_name, rng, rule_type = task["sheet"], task["range"], task["rule_type"]
        sheet = self._get_sheet(sheet_name)
        if sheet is None:
            self._log_audit("ERROR", f"Sheet '{sheet_name}' not found.", "conditional_format", task)
            return

        try:
            if rule_type == "color_scale":
                rule = ColorScaleRule(
                    start_type="min", start_color=task.get("min_color", "F8696B"),
                    mid_type="percentile", mid_value=50, mid_color=task.get("mid_color", "FFEB84"),
                    end_type="max", end_color=task.get("max_color", "63BE7B"),
                )
            elif rule_type == "data_bar":
                rule = DataBarRule(
                    start_type="min", end_type="max", color=task.get("color", "638EC6"), showValue=True
                )
            elif rule_type == "icon_set":
                rule = IconSetRule(
                    task.get("icon_set", "3TrafficLights1"),
                    "percent",
                    task.get("values", [0, 33, 67]),
                    showValue=None,
                )
            elif rule_type == "expression":
                if "formula" not in task:
                    raise ValueError("'expression' rule requires 'formula'.")
                fill_color = task.get("fill_color", "FFC7CE")
                rule = FormulaRule(
                    formula=[task["formula"].lstrip("=")],
                    fill=PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid"),
                    font=Font(color=task.get("font_color", "9C0006")),
                )
            else:
                self._log_audit("ERROR", f"Unsupported rule_type: {rule_type}", "conditional_format", task)
                return

            sheet.conditional_formatting.add(rng, rule)
            self._log_audit("SUCCESS", f"Conditional format '{rule_type}' applied to {rng} in '{sheet_name}'.", "conditional_format")
        except Exception as e:
            self._log_audit("ERROR", f"Failed to apply conditional format to {rng}: {e}", "conditional_format", task)

    def _process_summary_task(self, task):
        """สร้าง sheet สรุปด้วยสูตร Excel (SUM, AVERAGE, MAX, MIN, ...)"""
        src_name, dest_name = task["source_sheet"], task["sheet"]
        src = self._get_sheet(src_name)
        if src is None:
            self._log_audit("ERROR", f"Source sheet '{src_name}' not found.", "summary", task)
            return
        dest = self._get_sheet(dest_name, create=True)

        val_col, lbl_col = task["value_column"], task["label_column"]
        qty_col = task.get("quantity_column")
        first = task.get("start_row", 2)
        last = max(
            (r for r in range(first, src.max_row + 1) if src.cell(r, 1).value not in (None, "")),
            default=first,
        )
        s = self._quote(src_name)
        vals = f"{s}!${val_col}${first}:${val_col}${last}"
        lbls = f"{s}!${lbl_col}${first}:${lbl_col}${last}"

        rows = [
            ("Total Revenue", f"=SUM({vals})", "#,##0.00"),
            ("Average Revenue per Product", f"=AVERAGE({vals})", "#,##0.00"),
            ("Highest Revenue", f"=MAX({vals})", "#,##0.00"),
            ("Lowest Revenue", f"=MIN({vals})", "#,##0.00"),
            ("Top Product", f"=INDEX({lbls},MATCH(MAX({vals}),{vals},0))", "General"),
            ("Number of Products", f"=COUNTA({lbls})", "0"),
        ]
        if qty_col:
            qty = f"{s}!${qty_col}${first}:${qty_col}${last}"
            rows.insert(1, ("Total Units Sold", f"=SUM({qty})", "#,##0"))

        dest["A1"].value, dest["B1"].value = "Metric", "Value"
        for c in (dest["A1"], dest["B1"]):
            c.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        for i, (label, formula, fmt) in enumerate(rows, start=2):
            dest[f"A{i}"].value = label
            dest[f"A{i}"].font = Font(name=FONT_NAME)
            dest[f"B{i}"].value = formula
            dest[f"B{i}"].number_format = fmt
            dest[f"B{i}"].font = Font(name=FONT_NAME)
            dest[f"B{i}"].alignment = Alignment(horizontal="right")

        self._log_audit("SUCCESS", f"Summary sheet '{dest_name}' created with {len(rows)} metrics.", "summary")

    def _process_chart_task(self, task):
        sheet_name, chart_type = task["sheet"], task["chart_type"]
        sheet = self._get_sheet(sheet_name)
        if sheet is None:
            self._log_audit("ERROR", f"Sheet '{sheet_name}' not found.", "chart", task)
            return

        try:
            charts = {"bar": BarChart, "line": LineChart, "pie": PieChart}
            if chart_type not in charts:
                self._log_audit("ERROR", f"Unsupported chart type: {chart_type}", "chart", task)
                return
            chart = charts[chart_type]()

            min_col, min_row, max_col, max_row = range_boundaries(task["data_range"])
            data = Reference(sheet, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
            chart.add_data(data, titles_from_data=True)  # แถวแรกของ data_range = ชื่อ series

            if task.get("category_range"):
                c1, r1, c2, r2 = range_boundaries(task["category_range"])
                chart.set_categories(Reference(sheet, min_col=c1, min_row=r1, max_col=c2, max_row=r2))

            chart.title = task.get("title", "Chart")
            if chart_type != "pie":
                chart.x_axis.title = task.get("x_axis_title", "")
                chart.y_axis.title = task.get("y_axis_title", "")
                # openpyxl 3.1+ ซ่อนแกนโดยดีฟอลต์ ต้องตั้ง delete=False
                chart.x_axis.delete = False
                chart.y_axis.delete = False
            if chart_type == "bar":
                chart.legend = None
            chart.width, chart.height = task.get("width", 18), task.get("height", 9)

            sheet.add_chart(chart, task.get("top_left_cell", "G2"))
            self._log_audit("SUCCESS", f"{chart_type.capitalize()} chart '{task.get('title', 'Chart')}' added to '{sheet_name}'.", "chart")
        except Exception as e:
            self._log_audit("ERROR", f"Failed to create chart: {e}", "chart", task)

    def _process_data_validation_task(self, task):
        sheet_name, rng, vtype = task["sheet"], task["range"], task["validation_type"]
        sheet = self._get_sheet(sheet_name)
        if sheet is None:
            self._log_audit("ERROR", f"Sheet '{sheet_name}' not found.", "data_validation", task)
            return
        try:
            dv = DataValidation(
                type=vtype,  # whole | decimal | list | date | textLength
                operator=task.get("operator") if vtype != "list" else None,
                formula1=task.get("formula1"),
                formula2=task.get("formula2"),
                allow_blank=True,
                showErrorMessage=True,
                errorTitle=task.get("error_title", "Invalid input"),
                error=task.get("error_message", "Value does not meet the validation rule."),
            )
            sheet.add_data_validation(dv)
            dv.add(rng)
            self._log_audit("SUCCESS", f"Data validation ({vtype}) applied to {rng} in '{sheet_name}'.", "data_validation")
        except Exception as e:
            self._log_audit("ERROR", f"Failed to apply data validation: {e}", "data_validation", task)

    # ------------------------------------------------------------------ #
    # Output helpers
    # ------------------------------------------------------------------ #
    def _autofit_columns(self, sheet):
        for col_cells in sheet.columns:
            letter = get_column_letter(col_cells[0].column)
            longest = max((len(str(c.value)) for c in col_cells if c.value is not None and not str(c.value).startswith("=")), default=8)
            sheet.column_dimensions[letter].width = min(max(longest + 3, 12), 60)

    def _style_header_row(self, sheet):
        for cell in sheet[1]:
            if cell.value is not None and sheet.title != "Summary":
                cell.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

    def _add_audit_log_sheet(self):
        if "Audit Log" in self.workbook.sheetnames:
            del self.workbook["Audit Log"]
        ws = self.workbook.create_sheet("Audit Log")  # ไว้ท้ายสุด ไม่ให้บังหน้าแรก
        ws.append(["Timestamp", "Status", "Task Type", "Message", "Details"])
        for entry in self.audit_log:
            ws.append([
                entry["timestamp"], entry["status"], entry["task_type"], entry["message"],
                json.dumps(entry["details"], ensure_ascii=False) if entry["details"] else "",
            ])

        colors = {"SUCCESS": "C6EFCE", "WARNING": "FFEB9C", "ERROR": "FFC7CE", "CRITICAL": "FF9999"}
        for row in ws.iter_rows(min_row=2):
            fill = colors.get(row[1].value)
            for c in row:
                c.font = Font(name=FONT_NAME)
            if fill:
                row[1].fill = PatternFill(start_color=fill, end_color=fill, fill_type="solid")
        self._style_header_row(ws)
        self._autofit_columns(ws)
        ws.freeze_panes = "A2"

    # ------------------------------------------------------------------ #
    # Entry point
    # ------------------------------------------------------------------ #
    def run(self):
        if not self._load_workbook():
            return False

        success = True
        try:
            for i, task in enumerate(self.config["tasks"], start=1):
                task_type = task.get("type")
                logger.info(f"Executing Task {i}: type='{task_type}'")
                handler = self.handlers.get(task_type)
                if handler is None:
                    self._log_audit("WARNING", f"Unsupported task type: {task_type}", "task_execution", task)
                    continue
                try:
                    handler(task)
                except Exception as e:  # ให้ task เดียวพังได้โดยไม่ล้มทั้งระบบ
                    self._log_audit("ERROR", f"Task {i} ('{task_type}') crashed: {e}", task_type, task)

            # จัดรูปแบบ sheet ที่ agent สร้าง/แก้ไข
            for name in ("Raw Data", "Summary"):
                if name in self.workbook.sheetnames:
                    ws = self.workbook[name]
                    self._style_header_row(ws)
                    self._autofit_columns(ws)

            duration = (datetime.datetime.now() - self.start_time).total_seconds()
            self._log_audit("INFO", f"Agent execution completed in {duration:.2f} seconds.", "run_summary")
            self._add_audit_log_sheet()

            output_file = self.config["output_file"]
            self.workbook.save(output_file)
            logger.info(f"Final workbook saved to: {output_file}")
        except Exception as e:
            success = False
            logger.critical(f"Unhandled error during agent execution: {e}", exc_info=True)
        return success

    def has_errors(self):
        return any(e["status"] in ("ERROR", "CRITICAL") for e in self.audit_log)
