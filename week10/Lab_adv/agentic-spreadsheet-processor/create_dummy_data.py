# create_dummy_data.py
"""สร้างไฟล์ data/input_data.xlsx (sheet 'SalesData') สำหรับทดสอบ agent"""
import os

from openpyxl import Workbook
from openpyxl.styles import Font

SALES = [
    ("Laptop", 2, 1200.50),
    ("Mouse", 5, 25.00),
    ("Keyboard", 3, 75.99),
    ("Monitor", 1, 300.00),
    ("Projector", 1, 850.00),
    ("Webcam", 10, 45.00),
    ("Headset", 4, 99.99),
    ("SSD", 2, 180.00),
    ("Router", 1, 70.00),
]


def create_dummy_data(path=os.path.join("data", "input_data.xlsx")):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "SalesData"  # ต้องตรงกับ source_sheet ใน config

    ws.append(["Product Name", "Quantity", "Unit Price"])
    for row in SALES:
        ws.append(row)

    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in range(2, ws.max_row + 1):
        ws[f"C{r}"].number_format = "#,##0.00"
    for col, width in zip("ABC", (18, 12, 12)):
        ws.column_dimensions[col].width = width

    wb.save(path)
    print(f"Created {path} with {len(SALES)} rows.")


if __name__ == "__main__":
    create_dummy_data()
