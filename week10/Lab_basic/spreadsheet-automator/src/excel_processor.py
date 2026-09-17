import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

class ExcelProcessor:
    def __init__(self):
        pass

    def create_workbook(self, filename="new_workbook.xlsx"):
        wb = openpyxl.Workbook()
        if 'Sheet' in wb.sheetnames:
            del wb['Sheet']
        return wb

    def load_workbook(self, filename):
        try:
            wb = openpyxl.load_workbook(filename)
            return wb
        except Exception as e:
            print(f"Error loading workbook '{filename}': {e}")
            return None

    def save_workbook(self, workbook, filename):
        try:
            workbook.save(filename)
            return True
        except Exception as e:
            print(f"Error saving workbook: {e}")
            return False

    def process_sales_data(self, input_filename, output_filename):
        input_wb = self.load_workbook(input_filename)
        if not input_wb:
            return False
        
        input_sheet = input_wb.active
        output_wb = self.create_workbook()
        output_sheet = output_wb.create_sheet("Sales Report", 0)

        # Styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")
        total_row_font = Font(bold=True, color="000000")
        total_row_fill = PatternFill(start_color="FFEB3B", end_color="FFEB3B", fill_type="solid")
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        currency_format = '#,##0.00'

        headers = ["Product Name", "Quantity", "Unit Price", "Total Price"]
        for col_num, header in enumerate(headers, 1):
            cell = output_sheet.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')

        grand_total = 0
        output_row = 2

        for row_index, row in enumerate(input_sheet.iter_rows(min_row=2, values_only=True), 2):
            if not any(row):
                continue
            
            product_name = row[0] if len(row) > 0 else None
            quantity = row[1] if len(row) > 1 else 0
            unit_price = row[2] if len(row) > 2 else 0.0

            try:
                quantity = int(quantity)
                unit_price = float(unit_price)
            except (ValueError, TypeError):
                continue

            total_price = quantity * unit_price
            grand_total += total_price

            output_sheet.cell(row=output_row, column=1, value=product_name).border = thin_border
            output_sheet.cell(row=output_row, column=2, value=quantity).border = thin_border
            output_sheet.cell(row=output_row, column=3, value=unit_price).border = thin_border
            
            total_price_cell = output_sheet.cell(row=output_row, column=4, value=total_price)
            total_price_cell.number_format = currency_format
            total_price_cell.border = thin_border
            output_row += 1

        # Grand Total Row
        output_sheet.merge_cells(start_row=output_row, start_column=1, end_row=output_row, end_column=3)
        grand_total_label_cell = output_sheet.cell(row=output_row, column=1, value="Grand Total:")
        grand_total_label_cell.font = total_row_font
        grand_total_label_cell.fill = total_row_fill
        grand_total_label_cell.alignment = Alignment(horizontal='right', vertical='center')

        grand_total_value_cell = output_sheet.cell(row=output_row, column=4, value=grand_total)
        grand_total_value_cell.font = total_row_font
        grand_total_value_cell.fill = total_row_fill
        grand_total_value_cell.number_format = currency_format
        grand_total_value_cell.border = thin_border

        # Adjust column width
        for col_num in range(1, len(headers) + 1):
            max_length = 0
            column_letter = get_column_letter(col_num)
            for cell in output_sheet[column_letter]:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            output_sheet.column_dimensions[column_letter].width = max_length + 4

        return self.save_workbook(output_wb, output_filename)